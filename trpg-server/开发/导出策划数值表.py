# -*- coding: utf-8 -*-
"""导出项目内分散的玩法数值为 Excel 盘点表。

当前阶段是“汇总/审计”：Excel 不是运行时真相源，修改 Excel 不会自动回写代码或 JSON。
运行：在 trpg-server 下执行 `python 开发/导出策划数值表.py`。
输出：trpg-world/策划数值总表.xlsx
"""
from __future__ import annotations

import ast
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent.parent
SERVER = ROOT / "trpg-server"
WORLD = ROOT / "trpg-world"
CLIENT = ROOT / "trpg-client"
MAP = ROOT / "trpg-map"
OUT = WORLD / "策划数值总表.xlsx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GOLD = "FFF2CC"
LIGHT_RED = "FCE4D6"
LIGHT_GREEN = "E2F0D9"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="B7B7B7")


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8-sig"))


def py_constants(path: Path) -> dict[str, tuple[Any, int]]:
    """读取可由 ast.literal_eval 解析的模块级常量，返回 值+行号。"""
    tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
    out: dict[str, tuple[Any, int]] = {}
    for node in tree.body:
        if not isinstance(node, (ast.Assign, ast.AnnAssign)):
            continue
        targets = node.targets if isinstance(node, ast.Assign) else [node.target]
        value_node = node.value
        try:
            value = ast.literal_eval(value_node)
        except Exception:
            continue
        for target in targets:
            if isinstance(target, ast.Name):
                out[target.id] = (value, node.lineno)
            elif isinstance(target, (ast.Tuple, ast.List)) and isinstance(value, tuple):
                for elt, item in zip(target.elts, value):
                    if isinstance(elt, ast.Name):
                        out[elt.id] = (item, node.lineno)
    return out


def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def source(path: Path, line: int | None = None, key: str = "") -> str:
    s = rel(path)
    if line:
        s += f":{line}"
    if key:
        s += f" → {key}"
    return s


def text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, (dict, list, tuple)):
        return json.dumps(v, ensure_ascii=False, separators=(",", ":"))
    return str(v)


def add_table(wb: Workbook, title: str, headers: list[str], rows: list[list[Any]],
              note: str = "", tab_color: str = BLUE):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    start = 1
    if note:
        ws.cell(1, 1, note)
        ws.cell(1, 1).font = Font(italic=True, color="7F6000")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=LIGHT_GOLD)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        start = 3
    for col, h in enumerate(headers, 1):
        c = ws.cell(start, col, h)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    for r_idx, row in enumerate(rows, start + 1):
        for c_idx, value in enumerate(row, 1):
            c = ws.cell(r_idx, c_idx, value if isinstance(value, (int, float)) and not isinstance(value, bool) else text(value))
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            if r_idx % 2 == 0:
                c.fill = PatternFill("solid", fgColor="F7FBFE")
    if rows:
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
    ws.freeze_panes = f"A{start + 1}"
    for i, h in enumerate(headers, 1):
        max_len = len(h)
        for row in rows[:300]:
            if i - 1 < len(row):
                max_len = max(max_len, len(text(row[i - 1])))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 10), 42)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    # 说明
    ws = wb.create_sheet("说明")
    ws.sheet_properties.tabColor = "C65911"
    ws["A1"] = "TRPG 策划数值总表"
    ws["A1"].font = Font(size=20, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells("A1:F1")
    intro = [
        ("版本性质", "第一阶段：数值盘点与审计快照。当前游戏仍以 JSON / Python 常量为真相源。"),
        ("重要提醒", "直接修改此 Excel 暂时不会影响游戏；下一阶段需建立 Excel→JSON/代码配置的导入流程，或把运行时统一改读一个配置文件。"),
        ("包含范围", "战斗、技能树、养成、设施、时间/饥饿、移动/地形、天气、概率骰、世界推演、营业时间、经济参考及散落默认值。"),
        ("默认排除", "运行存档（游戏数据/*.json）、纯 UI 尺寸、模型 token/timeout、地图坐标和史实年份不作为平衡数值导入。"),
        ("生成脚本", rel(Path(__file__))),
        ("输出位置", rel(OUT)),
    ]
    for i, (k, v) in enumerate(intro, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    ws.column_dimensions["A"].width = 16
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20

    # 文件索引
    index_rows = [
        ["养成", "成长曲线/时长/战斗成长/技能点", "成长曲线、修行时长、战斗成长", "配置真相源", "热读", "trpg-server/配置/成长.json"],
        ["设施", "设施耗时/收益/buff", "设施数值", "配置真相源", "热读", "trpg-server/配置/facilities.json"],
        ["技能", "技能树节点/门槛/花费/招式数值", "技能树", "预制真相源", "重启/接口重读", "trpg-server/配置/技能树.json"],
        ["技能", "玩家已学招式运行表", "招式运行表", "运行投影", "运行时改写", "trpg-server/配置/招式表.json"],
        ["战斗", "核心/梯度/Buff/武器/位置/NPC推导/战术AI/思路评分", "战斗核心、Buff、NPC梯度、战术AI", "配置真相源", "热读", "trpg-server/配置/战斗数值.json"],
        ["战斗", "战斗公式与结算逻辑", "—", "代码逻辑（不保存策划系数）", "需重启", "trpg-server/tools/战斗/battle*.py"],
        ["时间", "精力/睡眠/饥饿消耗", "时间与生存", "配置真相源", "热读", "trpg-server/配置/时间影响.json"],
        ["时间", "刻/时辰/时钟倍率/战斗回合时长", "时间与生存", "代码默认", "需重启", "trpg-server/tools/核心/game_clock.py"],
        ["生存", "饥饿挡位阈值", "时间与生存", "代码常量", "需重启", "trpg-server/tools/核心/hunger.py"],
        ["移动", "基础步速/轻功/奔跑", "移动与地形", "配置真相源", "热读", "trpg-server/配置/移动.json"],
        ["移动", "地形与道路倍率", "移动与地形", "生成配置", "重导 walkable", "trpg-map/draw_tiles/export_walkable.py"],
        ["移动", "桥/码头放行半径", "移动与地形", "前端代码常量", "前端重载", "trpg-client/src/in-game/walkable.ts"],
        ["天气", "分区/极端天气权重/结构化规则影响", "天气", "配置真相源", "热读", "trpg-server/配置/天气.json"],
        ["世界推演", "每日顺利度权重", "世界推演与概率", "代码常量", "需重启", "trpg-server/tools/小模型/world_sim.py"],
        ["判定", "可能性骰档位与修正建议", "世界推演与概率", "规则文档", "热读", "trpg-world/主持人/可能性骰.md"],
        ["判定", "回合意外概率", "世界推演与概率", "代码逻辑", "需重启", "trpg-server/tools/大模型/accident.py"],
        ["营业", "地点营业时辰", "营业时间", "配置真相源", "需重建地图数据", "trpg-world/营业时间.json"],
        ["经济", "物价与收入参考", "经济参考", "规则文档", "热读", "trpg-world/主持人/世界.md"],
        ["提示词", "进食恢复/睡觉恢复等建议值", "散落数值审计", "提示文本", "需重启", "trpg-server/tools/大模型/registry.py"],
        ["运行存档", "当前生命/属性/金钱/背包等", "未纳入", "玩家存档，不是策划配置", "运行时变化", "trpg-server/游戏数据/*.json"],
        ["废弃", "日常/旅途内容骰", "未纳入", "已停用", "—", "trpg-world/废弃/*.md"],
    ]
    add_table(wb, "文件索引", ["系统", "数值内容", "Excel工作表", "当前地位", "生效方式", "源文件"], index_rows,
              "这里列出本次盘点发现的主要数值来源。")

    # 养成
    growth_path = SERVER / "配置" / "成长.json"
    growth = load_json(growth_path)
    curve_rows = []
    for kind, points in growth.get("曲线", {}).items():
        for threshold, coef in points:
            curve_rows.append([kind, threshold, coef, source(growth_path, key=f"曲线.{kind}")])
    add_table(wb, "成长曲线", ["类别", "当前值阈值", "收益系数", "来源"], curve_rows)

    duration_rows = [[ke, coef, round(coef / ke, 4), source(growth_path, key="时长曲线.表")]
                     for ke, coef in growth.get("时长曲线", {}).get("表", [])]
    add_table(wb, "修行时长", ["刻数", "收益系数", "每刻效率", "来源"], duration_rows,
              "每刻效率用于直观看出长练/短练的性价比。")

    battle_growth_rows = [[k, v, source(growth_path, key=f"战斗成长.{k}")]
                          for k, v in growth.get("战斗成长", {}).items() if not k.startswith("_")]
    for k, v in growth.get("技能点概率", {}).items():
        if not k.startswith("_"):
            battle_growth_rows.append([f"技能点概率·{k}", v, source(growth_path, key=f"技能点概率.{k}")])
    add_table(wb, "战斗成长", ["项目", "数值", "来源"], battle_growth_rows)

    # 设施
    fac_path = SERVER / "配置" / "facilities.json"
    facilities = load_json(fac_path)
    fac_rows = []
    for key, entry in facilities.items():
        if key.startswith("_") or not isinstance(entry, dict):
            continue
        for opt in entry.get("选项", []) or []:
            eff = opt.get("效果") or {}
            fac_rows.append([
                key, entry.get("名称", key), entry.get("类", ""), entry.get("行", ""), entry.get("耗时刻", ""),
                opt.get("标签", ""), opt.get("类型", ""), eff.get("类型", ""), eff.get("目标", ""),
                eff.get("点数", ""), eff.get("倍率", ""), eff.get("天数", ""), opt.get("介绍", ""),
                source(fac_path, key=f"{key}.选项.{opt.get('标签','')}")
            ])
    add_table(wb, "设施数值", ["配置键", "设施", "类别", "行", "默认耗时刻", "选项", "选项类型", "效果类型",
                               "成长目标", "基准点数", "Buff倍率", "持续天数", "说明", "来源"], fac_rows)

    # 技能树
    tree_path = SERVER / "配置" / "技能树.json"
    tree = load_json(tree_path)
    skill_rows = []
    for name, node in (tree.get("节点") or {}).items():
        req = node.get("要求") or {}
        skill_rows.append([
            name, node.get("五行", ""), node.get("层级", ""), "、".join(node.get("前置") or []),
            req.get("熟练度", ""), req.get("剑法", ""), node.get("花费", ""), node.get("内力", ""),
            node.get("射程", ""), node.get("范围", ""), node.get("威力", ""), node.get("连击", ""),
            "、".join(node.get("效果") or []), node.get("描述", ""), source(tree_path, key=f"节点.{name}")
        ])
    skill_rows.sort(key=lambda r: (r[1], r[2] if isinstance(r[2], (int, float)) else 99, r[0]))
    add_table(wb, "技能树", ["节点", "五行", "层级", "前置", "要求熟练度", "要求剑法", "技能点花费", "内力消耗",
                             "射程", "范围", "威力", "连击", "效果", "描述", "来源"], skill_rows,
              f"当前共 {len(skill_rows)} 个节点；此表是预制技能数值的主要编辑对象。")

    moves_path = SERVER / "配置" / "招式表.json"
    moves = load_json(moves_path)
    move_rows = []
    normal = moves.get("普通攻击") or {}
    if normal:
        move_rows.append(["普通攻击", normal.get("名称", "舞剑"), normal.get("五行"), normal.get("内力"), normal.get("射程"),
                          normal.get("范围"), normal.get("威力"), normal.get("连击"), "、".join(normal.get("效果") or []),
                          normal.get("描述", ""), source(moves_path, key="普通攻击")])
    for name, node in (moves.get("招式") or {}).items():
        move_rows.append(["已学招式", name, node.get("五行"), node.get("内力"), node.get("射程"), node.get("范围"),
                          node.get("威力"), node.get("连击"), "、".join(node.get("效果") or []), node.get("描述", ""),
                          source(moves_path, key=f"招式.{name}")])
    add_table(wb, "招式运行表", ["类型", "名称", "五行", "内力", "射程", "范围", "威力", "连击", "效果", "描述", "来源"], move_rows,
              "这是玩家已学技能的运行投影，会被技能树系统写入；不建议把它当唯一策划源。")

    # 战斗核心（唯一数值源：战斗数值.json）
    battle_cfg_path = SERVER / "配置" / "战斗数值.json"
    battle_cfg = load_json(battle_cfg_path)
    board = battle_cfg["棋盘"]
    hit = battle_cfg["命中"]
    wuxing = battle_cfg["五行"]
    defend = battle_cfg["防守"]
    charge = battle_cfg["蓄力"]
    retreat = battle_cfg["撤退"]
    damage = battle_cfg["伤害"]
    move_rule = battle_cfg["移动力"]
    core_rows = [
        ["棋盘宽", board["宽"], "格", source(battle_cfg_path, key="棋盘.宽")],
        ["棋盘高", board["高"], "格", source(battle_cfg_path, key="棋盘.高")],
        ["五行克制·优势", wuxing["优势倍率"], "倍率", source(battle_cfg_path, key="五行.优势倍率")],
        ["五行克制·劣势", wuxing["劣势倍率"], "倍率", source(battle_cfg_path, key="五行.劣势倍率")],
        ["流血叠层上限", battle_cfg["流血叠层上限"], "层", source(battle_cfg_path, key="流血叠层上限")],
        ["移动力", f"{move_rule['基础']} + floor(轻功/{move_rule['每档轻功']})", "格/回合", source(battle_cfg_path, key="移动力")],
        ["回合内力回复", f"round(内力上限×{battle_cfg['回合']['内力回复比例']})", "点/回合", source(battle_cfg_path, key="回合.内力回复比例")],
        ["防守回内力", defend["回复内力"], "点", source(battle_cfg_path, key="防守.回复内力")],
        ["防守减伤", defend["减伤"], "比例", source(battle_cfg_path, key="防守.减伤")],
        ["总减伤上限", defend["总减伤上限"], "比例", source(battle_cfg_path, key="防守.总减伤上限")],
        ["穿甲", f"目标减伤×{battle_cfg['穿甲']['减伤保留倍率']}", "比例", source(battle_cfg_path, key="穿甲")],
        ["水行蓄力", f"+{charge['每层增伤']:.0%}/层，最多{charge['叠层上限']}层", "伤害", source(battle_cfg_path, key="蓄力")],
        ["撤退成功线", retreat["成功线"], "d100最终值", source(battle_cfg_path, key="撤退.成功线")],
        ["撤退轻功加成上限", retreat["轻功加成上限"], "点", source(battle_cfg_path, key="撤退.轻功加成上限")],
        ["命中属性修正夹值", f"{hit['属性修正下限']}~{hit['属性修正上限']}", "点", source(battle_cfg_path, key="命中")],
        ["普通会心线", hit["会心线"], "d100最终值", source(battle_cfg_path, key="命中.会心线")],
        ["最低命中伤害", damage["最低伤害"], "点", source(battle_cfg_path, key="伤害.最低伤害")],
        ["五行熟练度伤害", f"+{wuxing['熟练度每点增伤']:.0%}/点", "倍率", source(battle_cfg_path, key="五行.熟练度每点增伤")],
    ]
    add_table(wb, "战斗核心", ["项目", "数值/公式", "单位", "来源"], core_rows)

    hit_rows = [
        ["未命中", f"<{hit['未命中线']}", 0.0],
        ["擦中", f"{hit['未命中线']}–{hit['标准命中线'] - 1}", hit["擦中伤害倍率"]],
        ["命中", f"{hit['标准命中线']}–会心线-1", hit["普通伤害倍率"]],
        ["会心", f"≥会心线（{hit['会心线']}）", hit["会心伤害倍率"]],
    ]
    for row in hit_rows:
        row.append(source(battle_cfg_path, key="命中"))
    add_table(wb, "命中档位", ["结果", "最终值", "伤害系数", "来源"], hit_rows)

    npc_formula = battle_cfg["NPC推导"]
    def npc_value(key, coef):
        rule = npc_formula[key]
        return round(rule["系数"] * coef + rule["基础"])
    tier_rows = [[tier, coef, npc_value("生命", coef), npc_value("内力", coef), npc_value("主兵器", coef),
                  npc_value("副剑拳", coef), npc_value("副暗器", coef), npc_value("轻功", coef),
                  source(battle_cfg_path, key=f"梯度系数.{tier}")]
                 for tier, coef in battle_cfg["梯度系数"].items()]
    add_table(wb, "NPC梯度", ["梯度", "系数", "默认生命", "默认内力", "主兵器", "副剑/拳", "副暗器", "轻功", "来源"], tier_rows)

    buff_rows = []
    for name, d in battle_cfg["Buff"].items():
        buff_rows.append([name, d.get("类型"), d.get("可叠层"), d.get("持续"), d.get("每回合生命", ""), d.get("命中", ""),
                          d.get("伤害", ""), d.get("受伤", ""), d.get("减伤", ""), d.get("治疗", ""),
                          d.get("技能伤害", ""), d.get("会心阈值", ""), d.get("跳过回合", ""), d.get("内力回复", ""),
                          source(battle_cfg_path, key=f"Buff.{name}")])
    add_table(wb, "Buff", ["名称", "类型", "可叠层", "持续回合", "每回合生命", "命中", "输出伤害", "承伤", "减伤", "治疗倍率",
                           "技能伤害", "会心阈值", "跳过回合", "允许回内", "来源"], buff_rows)

    weapon_rows = []
    for name, d in battle_cfg["武器类型"].items():
        weapon_rows.append([name, d.get("伤害"), d.get("命中效果") or "", d.get("会心效果") or "", source(battle_cfg_path, key=f"武器类型.{name}")])
    for name, d in battle_cfg["位置"].items():
        weapon_rows.append([f"位置·{name}", d["伤害"], f"命中{d['命中']:+}", "", source(battle_cfg_path, key=f"位置.{name}")])
    add_table(wb, "武器与位置", ["类型", "伤害倍率", "命中附加/修正", "会心附加", "来源"], weapon_rows)

    tactics = battle_cfg["战术AI"]
    unit_value = tactics["单位价值"]
    tactics_rows = [[key, value, source(battle_cfg_path, key=f"战术AI.{key}")]
                    for key, value in tactics.items() if key != "单位价值"]
    tactics_rows.extend([[f"单位价值·{key}", value, source(battle_cfg_path, key=f"战术AI.单位价值.{key}")]
                         for key, value in unit_value.items()])
    tactics_rows.extend([[f"思路评价·{name}", value, source(battle_cfg_path, key=f"思路评价.{name}")]
                         for name, value in battle_cfg["思路评价"].items()])
    add_table(wb, "战术AI", ["项目", "数值", "来源"], tactics_rows)

    # 时间、生存
    time_path = SERVER / "配置" / "时间影响.json"
    tcfg = load_json(time_path)
    clock_path = SERVER / "tools/核心/game_clock.py"
    cc = py_constants(clock_path)
    hunger_path = SERVER / "tools/核心/hunger.py"
    hc = py_constants(hunger_path)
    time_rows = [
        ["昼间每时辰精力消耗", tcfg["每时辰精力"]["昼"], "精力", source(time_path, key="每时辰精力.昼")],
        ["夜间每时辰精力消耗", tcfg["每时辰精力"]["夜"], "精力", source(time_path, key="每时辰精力.夜")],
        ["睡眠回满", tcfg["睡眠回满时辰"], "时辰", source(time_path, key="睡眠回满时辰")],
        ["每时辰饥饿消耗", tcfg["每时辰饥饿"], "饥饿", source(time_path, key="每时辰饥饿")],
        ["每时辰刻数", cc["KE_PER_SHICHEN"][0], "刻", source(clock_path, cc["KE_PER_SHICHEN"][1])],
        ["每刻游戏秒", cc.get("SECONDS_PER_KE", (900, 46))[0], "秒", source(clock_path, cc.get("SECONDS_PER_KE", (900, 46))[1])],
        ["时钟默认倍率", cc["DEFAULT_RATE"][0], "游戏秒/真实秒", source(clock_path, cc["DEFAULT_RATE"][1])],
        ["战斗每回合时间", cc["DEFAULT_BATTLE_SECONDS_PER_ROUND"][0], "游戏秒", source(clock_path, cc["DEFAULT_BATTLE_SECONDS_PER_ROUND"][1])],
    ]
    for threshold, label in hc["LEVELS"][0]:
        if threshold >= 0:
            time_rows.append([f"饥饿挡位·{label}", threshold, "下限", source(hunger_path, hc["LEVELS"][1])])
    add_table(wb, "时间与生存", ["项目", "数值", "单位", "来源"], time_rows)

    # 移动与地形
    movement_path = SERVER / "配置" / "移动.json"
    movement = load_json(movement_path)
    move_cfg_rows = [[k, v, "", source(movement_path, key=k)] for k, v in movement.items() if k != "说明"]
    terrain_path = MAP / "draw_tiles/export_walkable.py"
    terrain = py_constants(terrain_path)
    for (cat, tag), (name, mult) in terrain["TERRAIN_RULES"][0].items():
        move_cfg_rows.append([f"地形·{name}", mult, f"{cat}={tag}", source(terrain_path, terrain["TERRAIN_RULES"][1])])
    move_cfg_rows.append(["地形·山地", terrain["MOUNTAIN_RULE"][0][1], "ancient_kind=山", source(terrain_path, terrain["MOUNTAIN_RULE"][1])])
    for name, mult in terrain["ROAD_MULT"][0].items():
        move_cfg_rows.append([f"道路·{name}", mult, "道路倍率", source(terrain_path, terrain["ROAD_MULT"][1])])
    walkable_path = CLIENT / "src/in-game/walkable.ts"
    move_cfg_rows += [["桥放行半径", 100, "米", source(walkable_path, 28)], ["码头放行半径", 80, "米", source(walkable_path, 29)]]
    location_path = SERVER / "tools/大模型/location.py"
    move_cfg_rows += [["移动耗时提示·片刻", 300, "米以下", source(location_path, 48)],
                      ["移动耗时提示·约一刻", 1500, "米以下", source(location_path, 50)],
                      ["移动耗时提示·半时辰", 4000, "米以下", source(location_path, 52)]]
    add_table(wb, "移动与地形", ["项目", "数值", "条件/单位", "来源"], move_cfg_rows)

    # 天气（唯一数值源：天气.json；影响为结构化规则）
    weather_path = SERVER / "配置" / "天气.json"
    weather_cfg = load_json(weather_path)
    weather_rows = []
    for zone, choices in weather_cfg["极端天气权重"].items():
        for condition, weight in choices.items():
            weather_rows.append(["极端天气权重", zone, condition, "权重", weight, source(weather_path, key=f"极端天气权重.{zone}.{condition}")])
    for condition, effect in weather_cfg["影响"].items():
        for name, modifier in effect.get("判定修正", {}).items():
            weather_rows.append(["判定修正", "", condition, name, modifier, source(weather_path, key=f"影响.{condition}.判定修正.{name}")])
        for name, multiplier in effect.get("效果倍率", {}).items():
            weather_rows.append(["效果倍率", "", condition, name, multiplier, source(weather_path, key=f"影响.{condition}.效果倍率.{name}")])
        for field in ("效果禁用", "禁止行动", "强制行动", "建议行动", "环境标签"):
            for value in effect.get(field, []):
                weather_rows.append([field, "", condition, value, True, source(weather_path, key=f"影响.{condition}.{field}")])
        if not effect:
            weather_rows.append(["天气影响", "", condition, "无", "", source(weather_path, key=f"影响.{condition}")])
    for zone in weather_cfg["分区"]:
        rule = "；".join(f"{key}={value}" for key, value in zone.items() if key != "名称") or "兜底"
        weather_rows.append(["天气分区", zone["名称"], "", "范围", rule, source(weather_path, key=f"分区.{zone['名称']}")])
    add_table(wb, "天气", ["类别", "分区", "天气", "目标/规则", "数值", "来源"], weather_rows)

    # 世界推演与概率
    world_sim_path = SERVER / "tools/小模型/world_sim.py"
    wsc = py_constants(world_sim_path)
    prob_rows = [["世界推演顺利度", name, weight, "%", source(world_sim_path, wsc["SMOOTHNESS_WEIGHTS"][1])]
                 for name, weight in wsc["SMOOTHNESS_WEIGHTS"][0]]
    dice_path = WORLD / "主持人/可能性骰.md"
    for label, lo, hi in [("彻底失败", 0, 9), ("失败", 10, 29), ("有可能", 30, 59), ("小成功", 60, 89), ("彻底成功", 90, 100)]:
        prob_rows.append(["可能性骰结果", label, f"{lo}–{hi}", "d100最终值", source(dice_path, 31)])
    accident_path = SERVER / "tools/大模型/accident.py"
    prob_rows.append(["行动意外", "触发", "98–100（3%）", "d100", source(accident_path, 22)])
    add_table(wb, "世界推演与概率", ["系统", "结果/档位", "数值", "单位", "来源"], prob_rows)

    # 营业时间
    hours_path = WORLD / "营业时间.json"
    hours = load_json(hours_path)
    hour_rows = []
    for period, kinds in hours.get("时段", {}).items():
        for kind in kinds:
            hour_rows.append([kind, period, source(hours_path, key=f"时段.{period}")])
    add_table(wb, "营业时间", ["地点类型", "营业时段", "来源"], hour_rows)

    # 经济参考
    world_path = WORLD / "主持人/世界.md"
    econ_rows = [
        ["换算", "1两", 1000, "文", source(world_path, 163)], ["换算", "1贯", 1000, "文", source(world_path, 163)],
        ["物价", "炊饼", 10, "文", source(world_path, 171)], ["物价", "素面", 20, "文", source(world_path, 172)],
        ["物价", "肉面", 50, "文", source(world_path, 173)], ["物价", "茶一壶", 15, "文", source(world_path, 174)],
        ["物价", "客栈单间", 150, "文/晚", source(world_path, 175)], ["物价", "布鞋", 150, "文", source(world_path, 176)],
        ["物价", "普通铁刀", 1000, "文", source(world_path, 177)], ["物价", "上等剑", 10000, "文以上", source(world_path, 178)],
        ["物价", "雇驴", 100, "文/日", source(world_path, 179)], ["收入", "码头挑夫", 100, "文/日", source(world_path, 185)],
        ["收入", "体力劳工", "2–5", "贯/月", source(world_path, 186)], ["收入", "熟练工匠", "5–10", "贯/月", source(world_path, 187)],
        ["收入", "普通士兵", 1, "贯/月", source(world_path, 188)], ["收入", "高级商号账房", 10, "贯以上/月", source(world_path, 189)],
    ]
    add_table(wb, "经济参考", ["类别", "项目", "数值", "单位", "来源"], econ_rows)

    # 散落数值与审计问题
    scattered_rows = [
        ["已修", "意外概率", "注释与实现已统一为3%：d100掷出98–100触发。", source(accident_path, 3), source(accident_path, 22), "后续可迁入统一数值配置。"],
        ["已修", "技能点概率单一真相源", "代码后备常量已删除；缺失或非法时明确报错，正常运行只读取成长.json。", "trpg-server/tools/核心/skill_tree.py", source(growth_path, key="技能点概率"), "后续调整只需修改成长.json。"],
        ["已修", "技能点概率文档", "战斗系统.md 已改为只引用成长.json，不再抄写会过期的当前值。", "trpg-world/战斗系统.md:250", source(growth_path, key="技能点概率"), "保持配置为唯一数值源。"],
        ["已修", "饥饿消耗后备值", "正常运行以时间影响.json为准；hunger.py 注释与容错后备值已统一为当前每时辰5。", source(hunger_path, 14), source(time_path, key="每时辰饥饿"), "后续调整只需修改时间影响.json。"],
        ["已修", "移动参数单一真相源", "策划数值只保留在移动.json；后端严格校验，前端未取得配置时速度为0并暂缓移动。", source(movement_path), "trpg-server/tools/核心/movement.py；trpg-client/src/in-game/GameController.tsx", "后续调整只需修改移动.json。"],
        ["已修", "战斗数值单一真相源", "梯度、Buff、命中、武器、位置、NPC推导、战术权重与思路评分已迁入战斗数值.json。", source(battle_cfg_path), "trpg-server/tools/战斗/battle*.py", "后续调整只需修改战斗数值.json。"],
        ["已修", "天气影响已结构化", "判定修正、效果倍率/禁用、行动门禁和环境标签均由天气.json 提供；文本只作派生展示。", source(weather_path), "trpg-server/tools/核心/weather_system.py", "后续系统直接读取 基本信息.天气.影响。"],
        ["已修", "技能节点数量文档统一", f"技能树实际 {len(skill_rows)} 个节点；README、TODO 与交接现统一为该数量，并声明以技能树.json实际统计为准。", source(tree_path), "README.md；TODO.md；交接.md", "后续导表继续自动统计节点数。"],
        ["低", "招式表是运行投影", "技能树点亮后会写招式表；直接同时修改两处可能漂移。", source(tree_path), source(moves_path), "明确技能树为预制源、招式表为存档投影。"],
        ["待接入", "战斗回合时间", "默认60游戏秒/回合已定义，但战斗生命周期尚未接入时钟。", source(clock_path, cc["DEFAULT_BATTLE_SECONDS_PER_ROUND"][1]), "README/TODO", "接入后再做平衡验证。"],
        ["低", "设施时长曲线短练更高效", "2刻×1=0.5系数/刻；8刻×2=0.25/刻；16刻×4=0.25/刻。", source(growth_path, key="时长曲线"), "交接.md", "属于待拍板平衡项。"],
        ["低", "提示词中有隐藏恢复建议", "饱餐+40、小食+15、宴席+60；生命/精力工具描述还写睡觉+20，但实际睡眠由比例公式结算。", "trpg-server/tools/大模型/registry.py:180,222,238", source(time_path), "把食物与恢复量结构化。"],
    ]
    add_table(wb, "散落数值审计", ["优先级", "问题", "现状", "来源A", "来源B", "建议"], scattered_rows,
              "这张表是后续真正“统一数值源”时最重要的迁移清单。", tab_color="C65911")

    # 统一化路线
    plan_rows = [
        [1, "盘点", "生成当前 Excel（本次完成）", "不改游戏读取逻辑"],
        [2, "配置化", "战斗与天气数值已迁入独立 JSON；后续继续处理其他文本/代码数值", "代码只保留公式与校验"],
        [3, "Excel导入", "新增 Excel→JSON 导入器；校验重复键、概率和、前置节点、数值类型", "Excel 成为编辑入口"],
        [4, "自动导出", "JSON→Excel 继续保留，用于核对游戏实际生效值", "防手改 JSON 后漂移"],
        [5, "测试", "加入战斗伤害样例、成长周期、移动消耗、概率和等回归测试", "改数值后可快速验算"],
    ]
    add_table(wb, "统一化路线", ["阶段", "名称", "工作", "目标"], plan_rows,
              "建议不要让游戏直接读取 xlsx；更稳妥的是 Excel 作为编辑源，导出为版本可审查的 JSON。", tab_color="70AD47")

    # 全局样式
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = sheet.auto_filter.ref

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"已生成：{OUT}")
    print(f"工作表：{len(wb.sheetnames)}")
    print(f"技能节点：{len(skill_rows)}；设施选项：{len(fac_rows)}")


if __name__ == "__main__":
    main()
